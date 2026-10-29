from pathlib import Path
import sys,shutil
import numpy as np,pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
try:from pptx import Presentation
except ImportError:
 sys.path.insert(0,'/tmp/avery_slides_deps');from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
R=Path(__file__).resolve().parents[2];out=R/'results/presentation';figdir=out/'figures';figdir.mkdir(exist_ok=True)
order=['10','25','50','75','all'];cats=['Top 10%','Top 25%','Top 50%','Top 75%','All positive edges'];vals=[]
for m in ['DWI','fMRI']:
 d=pd.read_csv(R/f'results/final/{m}/GNN_BAG/{m}_GNN_threshold_comparison.csv')
 vals.append([float(d.loc[d.retain_positive_percent.astype(str).eq(v),'MAE_years'].iloc[0]) for v in order])
colors=['#007F86','#477AB7'];plt.rcParams.update({'font.family':'DejaVu Sans','font.size':15,'svg.fonttype':'none','pdf.fonttype':42})
fig,ax=plt.subplots(figsize=(12.5,6.1));y=np.arange(5);bh=.32
for k,(v,label,c) in enumerate(zip(vals,['DWI (916 CN scans)','fMRI (642 CN scans)'],colors)):
 yy=y+(k-.5)*bh;ax.barh(yy,v,height=bh*.90,color=c,label=label)
 for a,b in zip(yy,v):ax.text(b+.09,a,f'{b:.2f}',va='center',fontweight='bold',fontsize=15,color='#142B45')
ax.set_yticks(y,cats);ax.invert_yaxis();ax.set_xlim(0,11);ax.set_xlabel('CN out-of-fold MAE (years; lower is better)',labelpad=14);ax.grid(axis='x',alpha=.2);ax.set_axisbelow(True)
for side in ['top','right','left']:ax.spines[side].set_visible(False)
ax.tick_params(axis='y',length=0,pad=12);ax.legend(loc='upper center',bbox_to_anchor=(.5,-.16),ncol=2,frameon=False)
fig.subplots_adjust(left=.21,right=.96,top=.97,bottom=.25)
stem=figdir/'GNN_Threshold_Comparison_Chart_Only'
for ext in ['png','svg','pdf']:fig.savefig(str(stem)+'.'+ext,dpi=400,bbox_inches='tight',facecolor='white')
plt.close(fig)
p=out/'Connectome_Brain_Age_Results_English_Editable.pptx';r=Presentation(p)
def txt(s,x,y,w,h,value,size=16,color='142B45',bold=False):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True;pp=tf.paragraphs[0];pp.text=value;pp.font.name='Aptos';pp.font.size=Pt(size);pp.font.bold=bold;pp.font.color.rgb=RGBColor.from_string(color)
def rect(s,x,y,w,h,c):
 sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(c);sh.line.fill.background()
for s in r.slides:
 title=next((sh.text for sh in s.shapes if sh.has_text_frame and abs(sh.top-Inches(.65))<100),'')
 if title=='Ridge leads the full-cohort unimodal analyses':
  for sh in s.shapes:
   if sh.has_table:
    widths=[3.8,1.4,1.7,1.8,1.55,1.85]
    for col,w in zip(sh.table.columns,widths):col.width=Inches(w)
    for row in sh.table.rows:
     row.height=Inches(.72)
     for cell in row.cells:
      for pp in cell.text_frame.paragraphs:
       pp.font.size=Pt(18)
       for rr in pp.runs:rr.font.size=Pt(18)
 if title!='Graph sparsity has modality-specific effects':continue
 for sh in list(s.shapes):
  if Inches(1.5)<=sh.top<Inches(6.25):sh._element.getparent().remove(sh._element)
 left=3.25;width=8.7;top=1.85;rh=.69
 for tick in range(0,12,2):
  xx=left+width*tick/11;rect(s,xx,top-.03,.009,3.52,'DFE5EB');txt(s,xx-.15,5.43,.5,.3,str(tick),12,'526477')
 for i,cat in enumerate(cats):
  yy=top+i*rh;txt(s,.65,yy+.11,2.5,.45,cat,17)
  for k,v in enumerate(vals):
   by=yy+.04+k*.27;rect(s,left,by,width*v[i]/11,.22,colors[k].lstrip('#'));txt(s,left+width*v[i]/11+.08,by-.04,.75,.35,f'{v[i]:.2f}',15,bold=True)
 txt(s,8.0,5.78,4.5,.35,'MAE (years; lower is better)',13,'526477')
 for j,label in enumerate(['DWI (916 CN scans)','fMRI (642 CN scans)']):
  xx=.8+j*3.5;rect(s,xx,5.85,.23,.17,colors[j].lstrip('#'));txt(s,xx+.33,5.78,3.1,.4,label,14,'526477')
 s.notes_slide.notes_text_frame.text+='\nChart redrawn from the final selected-threshold comparison CSVs. Five thresholds, both modalities, raw CN OOF MAE. Editable native vector bars and text; standalone 400-dpi PNG and SVG are also saved.'
r.save(p);shutil.copyfile(p,out/'Connectome_Brain_Age_Results_English_Editable_with_Pipelines.pptx')
print('Updated model table and threshold chart; standalone exports:',stem)
