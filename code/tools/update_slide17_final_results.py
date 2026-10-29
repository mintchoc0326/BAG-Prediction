from pathlib import Path
import json,sys,shutil,hashlib
try:import pptx
except ImportError:sys.path.insert(0,'/tmp/avery_slides_deps')
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
R=Path(__file__).resolve().parents[2];OUT=R/'results/presentation'
p=OUT/'Connectome_Brain_Age_Results_English_Editable.pptx'
r=Presentation(p);s=r.slides[16]
assert any(sh.has_text_frame and sh.text=='PCA by dataset: fMRI' for sh in s.shapes)
f=R/'results/final/fMRI/harmonization_evaluation/fMRI_PCA_by_dataset_before_after.png'
run=json.loads((R/'results/final/fMRI/harmonization_run.json').read_text())
prov=json.loads((f.parent/'PCA_by_dataset_provenance.json').read_text())
for entry in prov['inputs']:assert hashlib.sha256(Path(entry['path']).read_bytes()).hexdigest()==entry['sha256']
images=[sh for sh in s.shapes if sh.shape_type==MSO_SHAPE_TYPE.PICTURE];assert len(images)==1
old=images[0];left,top,width,height=old.left,old.top,old.width,old.height
old._element.getparent().remove(old._element)
s.shapes.add_picture(str(f),left,top,width=width,height=height)
for sh in s.shapes:
 if sh.has_text_frame and sh.top>Inches(7) and sh.left<Inches(12):
  sh.text='Final fMRI run: '+run['run_id']+' | results/final/fMRI/harmonization_evaluation/'
  pp=sh.text_frame.paragraphs[0];pp.font.size=Pt(9);pp.font.color.rgb=RGBColor.from_string('526477')
s.notes_slide.notes_text_frame.text=json.dumps({'slide':17,'authoritative_source':run['source'],'run_id':run['run_id'],'figure':str(f),'scans':run['scans'],'features':run['features'],'cohort_counts':run['cohort_counts'],'PCA':prov['stages'],'method':'Before and after are independently fitted StandardScaler + PCA(2); fMRI median imputation. Same points/colors/counts as final exported results.','comparison':'Final features match previous run exactly; PCA distribution and percentages therefore remain unchanged.'},indent=2)
r.save(p);shutil.copyfile(p,OUT/'Connectome_Brain_Age_Results_English_Editable_with_Pipelines.pptx')
# Future full rebuilds should also read final PCA exports.
q=R/'code/tools/reconcile_slides.py';t=q.read_text().replace("R/f'results/checks/PCA/figures/{mod}_PCA_by_dataset_before_after.png'","R/f'results/final/{mod}/harmonization_evaluation/{mod}_PCA_by_dataset_before_after.png'")
q.write_text(t)
print('Updated slide 17 from verified final fMRI run:',run['run_id'])
