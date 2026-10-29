from pathlib import Path
import sys,shutil
import pandas as pd,numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
try:from pptx import Presentation
except ImportError:
 sys.path.insert(0,'/tmp/avery_slides_deps');from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
R=Path(__file__).resolve().parents[2];out=R/'results/presentation';f=out/'figures/Fusion_Model_Comparison_Chart_Only'
d=pd.read_csv(R/'results/final/Fusion/report_tables/Table4_models_pooled.csv')
labels={'DWI':'DWI Ridge','fMRI':'fMRI Ridge','EARLY_FUSION':'Ridge early fusion','WEIGHTED_EARLY_FUSION':'Weighted Ridge fusion','LATE_MEAN':'Ridge late mean','DWI_GNN':'DWI GNN','fMRI_GNN':'fMRI GNN','GNN_FUSION':'GNN fusion'}
names=[labels[m] for m in d.model];values=d.MAE_years.to_numpy();colors=['#007F86' if m=='WEIGHTED_EARLY_FUSION' else '#477AB7' for m in d.model]
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':16,'svg.fonttype':'none','pdf.fonttype':42})
fig,ax=plt.subplots(figsize=(12.5,6.6));y=np.arange(len(names));ax.barh(y,values,color=colors,height=.64)
for yy,v in zip(y,values):ax.text(v+.10,yy,f'{v:.2f}',va='center',fontsize=16,fontweight='bold',color='#142B45')
ax.set_yticks(y,names);ax.invert_yaxis();ax.set_xlim(0,11);ax.set_xlabel('CN out-of-fold MAE (years; lower is better)',labelpad=14)
ax.grid(axis='x',alpha=.18);ax.set_axisbelow(True);ax.tick_params(axis='y',length=0,pad=13)
for side in ['top','right','left']:ax.spines[side].set_visible(False)
fig.subplots_adjust(left=.30,right=.96,top=.98,bottom=.13)
for ext in ['png','svg','pdf']:fig.savefig(str(f)+'.'+ext,dpi=400,bbox_inches='tight',facecolor='white')
plt.close(fig)
p=out/'Connectome_Brain_Age_Results_English_Editable.pptx';r=Presentation(p);s=r.slides[38]
assert any(sh.has_text_frame and sh.text=='Weighted Ridge has the lowest paired-sample MAE' for sh in s.shapes)
for sh in list(s.shapes):
 if Inches(1.5)<=sh.top<Inches(6.25):sh._element.getparent().remove(sh._element)
def text(x,y,w,h,val,size=16,color='142B45',bold=False):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));pp=sh.text_frame.paragraphs[0];pp.text=val;pp.font.name='Aptos';pp.font.size=Pt(size);pp.font.bold=bold;pp.font.color.rgb=RGBColor.from_string(color)
def rect(x,y,w,h,c):
 sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(c);sh.line.fill.background()
left=4.;width=7.6;top=1.77;rh=.47
for tick in range(0,12,2):
 xx=left+width*tick/11;rect(xx,top-.04,.009,3.85,'DFE5EB');text(xx-.12,5.65,.5,.3,str(tick),12,'526477')
for i,(name,v,c) in enumerate(zip(names,values,colors)):
 yy=top+i*rh;text(.7,yy+.02,3.15,.37,name,17,bold=name=='Weighted Ridge fusion');rect(left,yy+.07,width*v/11,.28,c[1:]);text(left+width*v/11+.1,yy+.01,.8,.35,f'{v:.2f}',16,bold=True)
text(5.0,6.02,6.7,.4,'CN OOF MAE (years; lower is better)',15,'526477')
s.notes_slide.notes_text_frame.text+='\nRedrawn from final pooled paired-sample results; every model uses 387 CN scans / 330 participants. Native editable vector bars; values unchanged.'
r.save(p);shutil.copyfile(p,out/'Connectome_Brain_Age_Results_English_Editable_with_Pipelines.pptx')
print('Saved corrected page 39 and standalone chart:',f)
