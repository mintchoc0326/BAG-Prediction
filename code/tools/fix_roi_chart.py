from pathlib import Path
import sys,shutil
import pandas as pd,numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
try:from pptx import Presentation
except ImportError:
 sys.path.insert(0,'/tmp/avery_slides_deps');from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
R=Path(__file__).resolve().parents[2];out=R/'results/presentation';stem=out/'figures/Archived_GNN_ROI_Attribution_Chart_Only'
d=pd.read_csv(R/'results/explore/GNN_attribution/roi_importance.csv').head(7)
name_map={'Right_Pallidum':'Right pallidum','Left_Pallidum':'Left pallidum','ctx_lh_precuneus':'Left precuneus','ctx_rh_precuneus':'Right precuneus','ctx_rh_entorhinal':'Right entorhinal cortex','ctx_lh_fusiform':'Left fusiform cortex','ctx_lh_middletemporal':'Left middle temporal cortex'}
names=[name_map.get(v,v) for v in d.roi_name];values=d.mean_incident_absolute_years.to_numpy()
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':16,'svg.fonttype':'none','pdf.fonttype':42})
fig,ax=plt.subplots(figsize=(12.5,6));y=np.arange(7);ax.barh(y,values,height=.62,color='#007F86')
for yy,v in zip(y,values):ax.text(v+.035,yy,f'{v:.2f}',va='center',fontsize=16,fontweight='bold',color='#142B45')
ax.set_yticks(y,names);ax.invert_yaxis();ax.set_xlim(0,3.2);ax.set_xlabel('Mean incident absolute IG (years)',labelpad=15)
ax.grid(axis='x',alpha=.18);ax.set_axisbelow(True);ax.tick_params(axis='y',length=0,pad=12)
for side in ['top','right','left']:ax.spines[side].set_visible(False)
fig.subplots_adjust(left=.36,right=.97,top=.98,bottom=.15)
for ext in ['png','svg','pdf']:fig.savefig(str(stem)+'.'+ext,dpi=400,bbox_inches='tight',facecolor='white')
plt.close(fig)
p=out/'Connectome_Brain_Age_Results_English_Editable.pptx';r=Presentation(p)
def title(s):return next((sh.text for sh in s.shapes if sh.has_text_frame and abs(sh.top-Inches(.65))<100),'')
for s in r.slides:
 if title(s)=='fMRI carries larger attribution in this fusion model':
  for sh in s.shapes:
   if sh.has_table:
    for col,w in zip(sh.table.columns,[1.35,3.1,3.9,3.75]):col.width=Inches(w)
    for row in sh.table.rows:
     row.height=Inches(.8)
     for cell in row.cells:
      for pp in cell.text_frame.paragraphs:
       pp.font.size=Pt(18)
       for rr in pp.runs:rr.font.size=Pt(18)
 if title(s)!='Top ROI attributions emphasize fMRI regions':continue
 for sh in list(s.shapes):
  if Inches(1.5)<=sh.top<Inches(6.25):sh._element.getparent().remove(sh._element)
 def text(x,y,w,h,val,size=16,color='142B45',bold=False):
  sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));pp=sh.text_frame.paragraphs[0];pp.text=val;pp.font.name='Aptos';pp.font.size=Pt(size);pp.font.bold=bold;pp.font.color.rgb=RGBColor.from_string(color)
 def rect(x,y,w,h,c):
  sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(c);sh.line.fill.background()
 left=4.9;width=6.9;top=1.8;rh=.49
 for tick in np.arange(0,3.1,.5):
  xx=left+width*tick/3.2;rect(xx,top-.03,.009,3.55,'DFE5EB');text(xx-.15,5.45,.5,.3,f'{tick:g}',12,'526477')
 for i,(name,v) in enumerate(zip(names,values)):
  yy=top+i*rh;text(.65,yy+.02,4.15,.4,name,17);rect(left,yy+.08,width*v/3.2,.28,'007F86');text(left+width*v/3.2+.08,yy+.01,.7,.4,f'{v:.2f}',16,bold=True)
 text(5.4,5.95,6.4,.4,'Mean incident absolute IG (years)',16,'526477')
 s.notes_slide.notes_text_frame.text+='\nRedrawn clearly from archived reconstructed-model attribution; names expanded for readability without changing ROI identity. All seven displayed ROIs are fMRI. This is not a new exact-checkpoint attribution result.'
r.save(p);shutil.copyfile(p,out/'Connectome_Brain_Age_Results_English_Editable_with_Pipelines.pptx')
print('Updated modality table and ROI chart; exported',stem)
