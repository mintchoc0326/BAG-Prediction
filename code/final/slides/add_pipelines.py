from pathlib import Path
p=Path(__file__).resolve().with_name('build_slides.py')
exec(compile(p.read_text().split("path=OUT/")[0],str(p),'exec'))
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR
from lxml import etree
original=list(prs.slides._sldIdLst); inserts={}
def box(s,x,y,w,h,label):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string('EAF3F5');sh.line.color.rgb=RGBColor.from_string(teal)
 tf=sh.text_frame;tf.word_wrap=True;tf.vertical_anchor=MSO_ANCHOR.MIDDLE
 for i,line in enumerate(label.split('\n')):
  pp=tf.paragraphs[0] if i==0 else tf.add_paragraph();pp.text=line;pp.alignment=PP_ALIGN.CENTER;pp.font.name='Aptos';pp.font.size=Pt((14 if i==0 else 11) if h<1 else (17 if i==0 else 14));pp.font.bold=i==0;pp.font.color.rgb=RGBColor.from_string(navy)
 return sh
def arrow(s,x1,y1,x2,y2):
 sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2));sh.line.color.rgb=RGBColor.from_string(teal);sh.line.width=Pt(2)
 e=etree.SubElement(sh.line._get_or_add_ln(),'{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd');e.set('type','triangle')
def row(s,labels,y=2.2):
 n=len(labels);w=(12.1-(n-1)*.35)/n
 for i,label in enumerate(labels):
  x=.6+i*(w+.35);box(s,x,y,w,1.5,label)
  if i<n-1:arrow(s,x+w,y+.75,x+w+.35,y+.75)
def pipe(anchor,title,source,labels,detail,caption,y=2.2):
 s=slide(title,'EDITABLE PIPELINE',source=source);row(s,labels,y);text(s,.7,4.35,11.9,1.5,detail,21);call(s,caption)
 inserts.setdefault(anchor,[]).append(prs.slides._sldIdLst[-1]);return s
pipe(3,'Data preparation and quality-control pipeline','run_DWI_harmonization.py; run_fMRI_harmonization.py', ['Discover files\nCohort connectomes\n+ metadata','Standardize IDs\nParticipant / visit\nDataset / site','Matrix QC\n84 × 84; numeric\nFinite values','Match covariates\nAge, sex, diagnosis\nRecord exclusions','Retain inputs\nAligned metadata\n+ edge features'], 'Repeated visits retain a participant identifier for downstream grouped validation. Singleton technical batches are removed before ComBat.','Output: auditable eligible scans, aligned features and DATASET_SITE batches.')
pipe(4,'DWI harmonization pipeline','run_DWI_harmonization.py; DWI metadata/combat_covariates.csv',['DWI matrices\n4 cohorts\n84 × 84','Extract edges\nUpper triangle\n3,486 features','Transform\nlog1p\nEdge QC','Joint ComBat\nDATASET_SITE\nAge + sex*','Reconstruct\nInverse transform\nSymmetric matrices'],'*Diagnosis preservation is conditional in the DWI code; the saved run covariates determine the exact specification. Nonnegative DWI exports support graph construction.','Full-data harmonization is upstream of the current model CV; held-out age leakage remains possible.')
pipe(4,'fMRI harmonization pipeline','run_fMRI_harmonization.py; README.md',['fMRI matrices\nADNI / HABS / ADRC\n84 × 84','Extract edges\nUpper triangle\n3,486 correlations','Fisher-z\nz = arctanh(r)\nClipped correlations','Joint ComBat\nDATASET_SITE\nAge + sex + diagnosis','Reconstruct\nr = tanh(z)\nDiagonal = 1'],'Location and scale adjustment is performed jointly across eligible dataset/site batches. Harmonized matrices and a scan manifest are saved.','Output: 847 harmonized fMRI scans with traceable metadata.')
pipe(5,'Harmonization evaluation pipeline','dwi_evaluation.py; fMRI_evaluation.py; report_tables/Table3_harmonization.csv',['Aligned inputs\nBefore / after\nSame eligible target set','PCA views\nColor by dataset\nand site batch','Technical signal\nPredict dataset\nand site batch','Biological signal\nPredict age, sex\nand diagnosis','Compare\nMetrics + figures\nBefore → after'],'Technical targets use balanced accuracy. Biological diagnostics include classification balanced accuracy and age MAE / R². Eligible samples differ by target.','Assess reduced technical predictability together with retained biological information.')
pipe(9,'Single-modality Ridge pipeline','Ridge_GNN_BAG.py; Ridge_fMRI_BAG.py',['CN features\nDWI or fMRI\n3,486 edges','Outer split\n5 participant-\ngrouped folds','Inner tuning\n4 grouped folds\nScaler → Ridge alpha','Outer refit\nTraining-only scaler\nSelected alpha','Held-out output\nAge → BAG\nPooled OOF metrics'],'For each outer fold, tuning uses only outer-training participants. After OOF evaluation, fit a final CN model for descriptive full-cohort predictions.','All visits from a participant stay together; upstream ComBat remains outside this CV.')
pipe(9,'Single-modality GNN pipeline','GNN_DWI_BAG.py; GNN_fMRI_BAG.py',['Connectome\n84 nodes\n84-value profiles','Build graph\nPositive edges\nTest 10–100%','GINE × 3\nMessage passing\nEdge attributes','Readout\nGlobal mean pool\nAge regression','Grouped OOF\nCompare sparsity\nSelect lowest MAE'],'DWI node profiles are nonnegative; fMRI profiles retain signed correlations. Edge thresholding changes topology while node profiles remain full length.','Selected topology: DWI top 50%; fMRI top 10%. Selection and reporting use the same OOF comparison.')
pipe(12,'Cross-modality pairing and training selection','prepare_DWI_fMRI_fusion_metadata.py; metadata/fusion_metadata_audit_summary.json',['Join modalities\nDataset + normalized\nparticipant / visit','Audit agreement\nAge / sex / diagnosis\nPreserve conflicts','Apply CN policy\nExisting modality\npolicy intersection','Filter training\nAge difference\n≤ 0.01 year','Freeze folds\n387 CN scans\n330 participants'],'847 paired scans are available. The DWI HABS all-CN override is retained; explicit fMRI CN is required. Age filtering removes most HABS pairs from training.','The same participant-grouped outer folds are reused across the paired model comparisons.')
s=slide('Ridge fusion pipelines: early, weighted and late','EDITABLE PIPELINE',source='Ridge_DWI_fMRI_fusion.py; Ridge_DWI_fMRI_weighted_fusion.py; fusion_transformers.py')
for y,label,steps in [(1.7,'EARLY FUSION',['DWI + fMRI edges','Scale columns\nTraining split only','Concatenate\n6,972 features','Tune alpha\nGrouped inner CV','Refit → OOF age']), (3.05,'WEIGHTED EARLY',['DWI + fMRI edges','Scale columns\nTraining split only','[Z_DWI, w·Z_fMRI]','Tune w + alpha\nGrouped inner CV','Refit → OOF age']), (4.4,'LATE MEAN',['DWI / fMRI edges','Separate scalers\n+ Ridge tuning','DWI age\nfMRI age','Average ages\n(DWI + fMRI) / 2','OOF age → BAG'])]:
 text(s,.62,y,12,.3,label,12,teal,True)
 n=len(steps);w=(12.1-4*.35)/5
 for j,st in enumerate(steps):
  x=.6+j*(w+.35);box(s,x,y+.36,w,.84,st)
  if j<4:arrow(s,x+w,y+.78,x+w+.35,y+.78)
call(s,'Apply block weights after scaling; scaling again would cancel the weighting.',y=6.45);inserts.setdefault(14,[]).append(prs.slides._sldIdLst[-1])
s=slide('GNN fusion pipeline: two modality-specific encoders','EDITABLE PIPELINE',source='GNN_DWI_fMRI_fusion.py')
for y,mod,prep in [(1.9,'DWI','Nonnegative edges\nlog1p'),(3.9,'fMRI','Signed profiles\nPositive topology')]:
 for x,w,t in [(.6,2.35,mod+' connectome\n'+prep),(3.3,2.4,'Graph + scaling\nTop 50% positive edges\nTraining node statistics'),(6.05,2.35,'GINE encoder × 2\nHidden size = 32\nMean + max pooling')]:box(s,x,y,w,1.5,t)
 arrow(s,2.95,y+.75,3.3,y+.75);arrow(s,5.7,y+.75,6.05,y+.75)
box(s,9.05,2.9,1.7,1.5,'Concatenate\nEmbeddings');box(s,11.1,2.9,1.65,1.5,'Age head\nMLP\nOOF age')
arrow(s,8.4,2.65,9.05,3.4);arrow(s,8.4,4.65,9.05,3.9);arrow(s,10.75,3.65,11.1,3.65)
call(s,'Inner grouped validation selects epochs → fresh outer-training refit → held-out prediction.')
s.notes_slide.notes_text_frame.text+='\nEdge attributes are normalized by the per-graph positive maximum; node and age statistics are fitted on training data. Single-modality paired GNN comparators use the corresponding one-encoder branch.'
inserts.setdefault(14,[]).append(prs.slides._sldIdLst[-1])
pipe(11,'BAG and age-bias correction pipeline','Ridge CV summaries; GNN threshold summaries',['OOF predictions\nCN held-out ages\n+ chronological age','Compute BAG\nPredicted age\n− actual age','Fit age trend\nCN OOF BAG\n= intercept + slope·age','Calculate cBAG\nBAG − fitted\nage-related trend','Summarize\nAge plots / cohort\nDescriptive outputs'],'The CN OOF reference defines the correction. Full-cohort CN predictions from the final model are in-sample and are not an independent performance test.','Independent or cross-fitted calibration is needed for stronger downstream inference.')
pipe(17,'Interpretability pipeline: from OOF models to attribution','GNN_BAG_interpretability.py; GNN_interpretability/METHODS.txt',['Reconstruct folds\nRecorded epochs\nHeld-out CN graphs','Prediction check\nCompare original\nvs reconstructed','Integrated gradients\nNode + edge inputs\nFixed topology','Quality checks\nCompleteness\nAdaptive integration','Aggregate\nParticipant → edge\nROI / modality'],'Use training-mean node baselines with the observed diagonal fixed and zero edge attributes. Average signed attribution within participant before absolute-value aggregation.','Reconstruction difference reaches 13.18 years; report these as reconstructed-model explanations.')
pipe(18,'Interpretability summaries and robustness checks','GNN_interpretability/*.csv; analysis_summary.json',['Connection scores\nCombine node / edge\nAttribution channels','ROI scores\nHalf of |connection IG|\nto each endpoint','Modality scores\nSigned + absolute\nParticipant summaries','Stability checks\nFold ranks\nTop-k frequency','Reference occlusion\nReplace one modality\nCompare error'],'Six scans fail the 0.05-year IG completeness tolerance. Reference occlusion retains graph topology and may be outside the observed data distribution.','Model reliance and rank stability are descriptive; neither establishes causal or clinical relevance.')
# Place pipeline slides next to their corresponding methods/results, preserving original slides.
allids=list(prs.slides._sldIdLst)
for el in allids:prs.slides._sldIdLst.remove(el)
for i,el in enumerate(original,1):
 prs.slides._sldIdLst.append(el)
 for extra in inserts.get(i,[]):prs.slides._sldIdLst.append(extra)
for n,s in enumerate(prs.slides,1):
 for sh in s.shapes:
  if sh.has_text_frame and sh.left>Inches(12) and sh.top>Inches(7):sh.text=f'{n:02}';sh.text_frame.paragraphs[0].font.size=Pt(10)
path=OUT/'Connectome_Brain_Age_Results_English_Editable_with_Pipelines.pptx';prs.save(path)
check=Presentation(path)
assert len(check.slides)==len(original)+sum(map(len,inserts.values()))
for n,s in enumerate(check.slides,1):
 for sh in s.shapes:
  assert sh.left>=0 and sh.top>=0 and sh.left+sh.width<=check.slide_width+10000 and sh.top+sh.height<=check.slide_height+10000,(n,sh.name)
print('Saved',path,'slides',len(check.slides),'pipeline slides',sum(map(len,inserts.values())))
